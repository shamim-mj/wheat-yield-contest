"""
Kentucky Wheat Yield Contest - Digital Entry Form
University of Kentucky Cooperative Extension
"""

import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import datetime, requests, json, time, base64
from pathlib import Path

# ═════════════════════════════════════════════════════════
#  OWNER CONFIG
# ═════════════════════════════════════════════════════════
FORMSUBMIT_EMAIL  = "shamim.one@outlook.com"
CC_EMAIL          = "mshamim11@uky.edu"
CONTACT_EMAIL     = "chad.lee@uky.edu"
EXCEL_FILENAME    = "wheat_contest_entries.xlsx"
EXCEL_FILE        = f"/tmp/{EXCEL_FILENAME}"
CURRENT_YEAR      = datetime.date.today().year
# ═════════════════════════════════════════════════════════

KY_COUNTIES = sorted([
    "Adair","Allen","Anderson","Ballard","Barren","Bath","Bell","Boone",
    "Bourbon","Boyd","Boyle","Bracken","Breathitt","Breckinridge","Bullitt",
    "Butler","Caldwell","Calloway","Campbell","Carlisle","Carroll","Carter",
    "Casey","Christian","Clark","Clay","Clinton","Crittenden","Cumberland",
    "Daviess","Edmonson","Elliott","Estill","Fayette","Fleming","Floyd",
    "Franklin","Fulton","Gallatin","Garrard","Grant","Graves","Grayson",
    "Green","Greenup","Hancock","Hardin","Harlan","Harrison","Hart",
    "Henderson","Henry","Hickman","Hopkins","Jackson","Jefferson","Jessamine",
    "Johnson","Kenton","Knott","Knox","LaRue","Laurel","Lawrence","Lee",
    "Leslie","Letcher","Lewis","Lincoln","Livingston","Logan","Lyon",
    "Madison","Magoffin","Marion","Marshall","Martin","Mason","McCracken",
    "McCreary","McLean","Meade","Menifee","Mercer","Metcalfe","Monroe",
    "Montgomery","Morgan","Muhlenberg","Nelson","Nicholas","Ohio","Oldham",
    "Owen","Owsley","Pendleton","Perry","Pike","Powell","Pulaski",
    "Robertson","Rockcastle","Rowan","Russell","Scott","Shelby","Simpson",
    "Spencer","Taylor","Todd","Trigg","Trimble","Union","Warren","Washington",
    "Wayne","Webster","Whitley","Wolfe","Woodford"
])

COUNTY_AREA = {
    "Ballard":1,"Calloway":1,"Carlisle":1,"Fulton":1,"Graves":1,"Hickman":1,
    "Marshall":1,"McCracken":1,"Caldwell":1,"Christian":1,"Crittenden":1,
    "Hopkins":1,"Lyon":1,"Todd":1,"Trigg":1,
    "Daviess":2,"Hancock":2,"Henderson":2,"McLean":2,"Muhlenberg":2,
    "Ohio":2,"Union":2,"Webster":2,
    "Adair":3,"Allen":3,"Barren":3,"Butler":3,"Edmonson":3,"Hart":3,
    "Logan":3,"Metcalfe":3,"Monroe":3,"Simpson":3,"Warren":3,
}
HEADER_FILL = "2E4057"

COLUMNS = [
    "Entry_ID","Submission_Date","County","Area",
    "Producer_Name","Producer_Email","Producer_Phone","Producer_Mobile",
    "Producer_Address","Producer_Town","Producer_Zip","Profession",
    "Supervisor_Name","Supervisor_Phone","Supervisor_Signature_Date",
    "Division","Previous_Crop","Planting_Date","Harvest_Date","Wheat_Variety",
    "Row_Width_inches","Seeding_Rate",
    "Fall_N_lbA","Fall_P2O5_lbA","Fall_K2O_lbA","Fall_Other_Fertilizer",
    "Winter_Spring_N1_Date","Winter_Spring_N1_lbA",
    "Winter_Spring_N2_Date","Winter_Spring_N2_lbA",
    "Manure_Used","Manure_Type","Manure_TonsA","Manure_Date",
    "Growth_Regulator","Fall_Pest_Products","Spring_Pest_Products",
    "Heading_Flowering_Pest","Biologicals_Other","Tillage_Used",
    "Harvest_Length_ft","Harvest_Width_ft","Harvest_Area_ft2","Harvest_Acres",
    "Grain_Moisture_1","Grain_Moisture_2","Grain_Moisture_3","Grain_Moisture_Avg",
    "Test_Weight_lbbu","Grain_Weight_lbs","Official_Yield_BuAcre","Agent_Notes",
    "Scale_Ticket_Photo",
]

SECTION_SPANS = [
    ("Producer Info",          1, 12, "1F4E79"),
    ("Supervisor Info",       13, 15, "2E5FA3"),
    ("Agronomic Data",        16, 22, "375623"),
    ("Fertilizer",            23, 34, "7B3F00"),
    ("Pest Management",       35, 40, "6B2737"),
    ("Harvest Area",          41, 44, "4A235A"),
    ("Grain Characteristics", 45, 49, "7E5109"),
    ("Yield Calculation",     50, 51, "1A5276"),
    ("Notes & Photo",         52, 53, "555555"),
]

# ─────────────────────────────────────────────────────────
# GOOGLE SHEETS
# ─────────────────────────────────────────────────────────

def _gdrive_secrets() -> dict:
    try:
        return dict(st.secrets.get("gdrive", {}))
    except Exception:
        return {}

def _get_sheets_token(client_email: str, private_key: str) -> str | None:
    try:
        now = int(time.time())
        header  = {"alg": "RS256", "typ": "JWT"}
        payload = {
            "iss":   client_email,
            "scope": "https://www.googleapis.com/auth/spreadsheets",
            "aud":   "https://oauth2.googleapis.com/token",
            "iat":   now, "exp": now + 3600,
        }
        def b64(data: bytes) -> str:
            return base64.urlsafe_b64encode(data).rstrip(b"=").decode()
        h = b64(json.dumps(header).encode())
        p = b64(json.dumps(payload).encode())
        msg = f"{h}.{p}".encode()
        from cryptography.hazmat.primitives import hashes, serialization
        from cryptography.hazmat.primitives.asymmetric import padding as asym_padding
        key_str = private_key.replace("\\n", "\n")
        pk  = serialization.load_pem_private_key(key_str.encode(), password=None)
        sig = pk.sign(msg, asym_padding.PKCS1v15(), hashes.SHA256())
        jwt = f"{h}.{p}.{b64(sig)}"
        resp = requests.post(
            "https://oauth2.googleapis.com/token",
            data={"grant_type": "urn:ietf:params:oauth:grant-type:jwt-bearer",
                  "assertion": jwt}, timeout=15)
        data = resp.json()
        if "access_token" in data:
            return data["access_token"]
        st.session_state["_gd_error"] = f"Token error: {data}"
        return None
    except Exception as e:
        st.session_state["_gd_error"] = f"JWT error: {e}"
        return None

def append_entry_to_sheet(data: dict, entry_id: int) -> bool:
    cfg          = _gdrive_secrets()
    client_email = cfg.get("client_email", "")
    private_key  = cfg.get("private_key", "")
    sheet_id     = cfg.get("sheet_id", "")
    if not all([client_email, private_key, sheet_id]):
        st.session_state["_gd_error"] = "Missing secrets"
        return False
    token = _get_sheets_token(client_email, private_key)
    if not token:
        return False
    try:
        r = data.get("_moisture_list", [])
        row = [
            entry_id, data.get("Submission_Date",""), data.get("County",""),
            COUNTY_AREA.get(data.get("County",""), 4),
            data.get("Producer_Name",""), data.get("Producer_Email",""),
            data.get("Producer_Phone",""), data.get("Producer_Mobile",""),
            data.get("Producer_Address",""), data.get("Producer_Town",""),
            data.get("Producer_Zip",""), data.get("Profession",""),
            data.get("Supervisor_Name",""), data.get("Supervisor_Phone",""),
            data.get("Supervisor_Signature_Date",""),
            data.get("Division",""), data.get("Previous_Crop",""),
            data.get("Planting_Date",""), data.get("Harvest_Date",""),
            data.get("Wheat_Variety",""),
            data.get("Row_Width_inches",""), data.get("Seeding_Rate",""),
            data.get("Fall_N_lbA",""), data.get("Fall_P2O5_lbA",""),
            data.get("Fall_K2O_lbA",""), data.get("Fall_Other_Fertilizer",""),
            data.get("Winter_Spring_N1_Date",""), data.get("Winter_Spring_N1_lbA",""),
            data.get("Winter_Spring_N2_Date",""), data.get("Winter_Spring_N2_lbA",""),
            data.get("Manure_Used",""), data.get("Manure_Type",""),
            data.get("Manure_TonsA",""), data.get("Manure_Date",""),
            data.get("Growth_Regulator",""), data.get("Fall_Pest_Products",""),
            data.get("Spring_Pest_Products",""), data.get("Heading_Flowering_Pest",""),
            data.get("Biologicals_Other",""), data.get("Tillage_Used",""),
            data.get("Harvest_Length_ft",""), data.get("Harvest_Width_ft",""),
            data.get("Harvest_Area_ft2",""), data.get("Harvest_Acres",""),
            r[0] if len(r)>0 else "", r[1] if len(r)>1 else "",
            r[2] if len(r)>2 else "", data.get("Grain_Moisture_Avg",""),
            data.get("Test_Weight_lbbu",""), data.get("Grain_Weight_lbs",""),
            data.get("Official_Yield_BuAcre",""), data.get("Agent_Notes",""),
            data.get("Scale_Ticket_Photo","[no photo]"),
        ]
        headers  = {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}
        base_url = f"https://sheets.googleapis.com/v4/spreadsheets/{sheet_id}"
        if entry_id == 1:
            hdr = [
                "Entry ID","Submission Date","County","Area",
                "Producer Name","Producer Email","Phone","Mobile",
                "Address","Town","Zip","Profession",
                "Supervisor Name","Supervisor Phone","Supervisor Sign Date",
                "Division","Previous Crop","Planting Date","Harvest Date","Wheat Variety",
                "Row Width (in)","Seeding Rate",
                "Fall N","Fall P2O5","Fall K2O","Fall Other",
                "Spring N1 Date","Spring N1 lb/A","Spring N2 Date","Spring N2 lb/A",
                "Manure","Manure Type","Manure T/A","Manure Date",
                "Growth Reg","Fall Pest","Spring Pest","Head/Flower Pest",
                "Biologicals","Tillage",
                "Length ft","Width ft","Area ft2","Acres",
                "Moisture 1","Moisture 2","Moisture 3","Moisture Avg",
                "Test Wt","Grain Wt lbs","Official Yield Bu/A","Agent Notes",
                "Scale Ticket Photo",
            ]
            requests.post(f"{base_url}/values/Sheet1!A1:append",
                          headers=headers, params={"valueInputOption":"RAW"},
                          json={"values":[hdr]}, timeout=15)
        resp = requests.post(
            f"{base_url}/values/Sheet1!A1:append",
            headers=headers,
            params={"valueInputOption":"RAW","insertDataOption":"INSERT_ROWS"},
            json={"values":[row]}, timeout=15)
        if resp.status_code == 200:
            return True
        st.session_state["_gd_error"] = f"HTTP {resp.status_code}: {resp.text[:300]}"
        return False
    except Exception as e:
        st.session_state["_gd_error"] = str(e)
        return False

# ─────────────────────────────────────────────────────────
# FORMSUBMIT — server-side, never fires on page load
# ─────────────────────────────────────────────────────────

def send_formsubmit_email(data: dict, entry_id: int, subject: str) -> bool:
    area = COUNTY_AREA.get(data.get("County",""), 4)
    r    = data.get("_moisture_list", [])
    readings_str = ", ".join(f"{x}%" for x in r) if r else "N/A"
    body = f"""KY WHEAT YIELD CONTEST — ENTRY #{entry_id}
{'='*50}
PRODUCER
  County:       {data.get('County','')} (Area {area})
  Producer:     {data.get('Producer_Name','')}
  Email:        {data.get('Producer_Email','')}
  Phone:        {data.get('Producer_Phone','')}
  Supervisor:   {data.get('Supervisor_Name','')}  Ph: {data.get('Supervisor_Phone','')}

AGRONOMIC
  Division:     {data.get('Division','')}
  Variety:      {data.get('Wheat_Variety','')}
  Planted:      {data.get('Planting_Date','')}
  Harvested:    {data.get('Harvest_Date','')}

HARVEST AREA
  {data.get('Harvest_Length_ft',0)} ft x {data.get('Harvest_Width_ft',0)} ft = {data.get('Harvest_Acres',0):.2f} acres

GRAIN
  Moisture: {readings_str}  →  Avg {data.get('Grain_Moisture_Avg',0):.1f}%
  Grain Weight: {data.get('Grain_Weight_lbs',0):,.0f} lbs

OFFICIAL YIELD: {data.get('Official_Yield_BuAcre',0):.2f} bu/acre

Submitted: {data.get('Submission_Date','')}"""
    try:
        resp = requests.post(
            f"https://formsubmit.co/ajax/{FORMSUBMIT_EMAIL}",
            headers={"Content-Type": "application/json", "Accept": "application/json"},
            json={"subject": subject, "cc": CC_EMAIL,
                  "_captcha": "false", "message": body},
            timeout=15)
        return resp.status_code == 200
    except Exception as e:
        st.session_state["_email_error"] = str(e)
        return False

# ─────────────────────────────────────────────────────────
# upload weight scale to google drive
# ─────────────────────────────────────────────────────────


def upload_photo_to_gdrive(photo_bytes: bytes, filename: str,
                            entry_id: int, county: str) -> str:
    """
    Upload scale ticket photo to Google Drive as an actual image file.
    Returns a shareable view URL, or empty string on failure.
    Uses the same service account token as Sheets.
    """
    cfg          = _gdrive_secrets()
    client_email = cfg.get("client_email", "")
    private_key  = cfg.get("private_key", "")
    folder_id    = cfg.get("photo_folder_id", cfg.get("sheet_id", ""))

    if not all([client_email, private_key]):
        return ""

    token = _get_sheets_token(client_email, private_key)
    if not token:
        return ""

    try:
        # Determine MIME type from filename
        ext  = filename.lower().split(".")[-1]
        mime = {"jpg": "image/jpeg", "jpeg": "image/jpeg",
                "png": "image/png",  "heic": "image/heic",
                "webp": "image/webp"}.get(ext, "image/jpeg")

        # Rename file to include entry info
        safe_name = f"Entry_{entry_id}_{county}_{filename}"

        # If photo_folder_id is set, upload there; otherwise upload to root
        parents = [folder_id] if cfg.get("photo_folder_id") else []
        metadata = json.dumps({"name": safe_name, "parents": parents}
                              if parents else {"name": safe_name})

        boundary = "====photo_upload===="
        body = (
            f"--{boundary}\r\n"
            f"Content-Type: application/json; charset=UTF-8\r\n\r\n"
            f"{metadata}\r\n"
            f"--{boundary}\r\n"
            f"Content-Type: {mime}\r\n\r\n"
        ).encode("utf-8") + photo_bytes + f"\r\n--{boundary}--".encode("utf-8")

        resp = requests.post(
            "https://www.googleapis.com/upload/drive/v3/files",
            headers={
                "Authorization": f"Bearer {token}",
                "Content-Type": f"multipart/related; boundary={boundary}",
            },
            params={"uploadType": "multipart", "fields": "id"},
            data=body,
            timeout=30,
        )

        if resp.status_code in (200, 201):
            file_id = resp.json().get("id", "")
            # Make file publicly viewable so the link works for anyone
            requests.post(
                f"https://www.googleapis.com/drive/v3/files/{file_id}/permissions",
                headers={"Authorization": f"Bearer {token}",
                         "Content-Type": "application/json"},
                json={"role": "reader", "type": "anyone"},
                timeout=10,
            )
            return f"https://drive.google.com/file/d/{file_id}/view"
        return ""
    except Exception:
        return ""













# ─────────────────────────────────────────────────────────
# SESSION STATE
# ─────────────────────────────────────────────────────────

def _today():
    return datetime.date.today()

def _blank_defaults():
    today = _today()
    return {
        "county": "— Select —",
        "producer_name": "", "producer_email": "",
        "producer_phone": "", "producer_mobile": "",
        "producer_address": "", "producer_town": "",
        "producer_zip": "", "profession": "",
        "supervisor_name": "", "supervisor_phone": "",
        "supervisor_sig_date": today,
        "division": "Division I - Tillage (conv./min.)",
        "previous_crop": "Corn",
        "planting_date": datetime.date(today.year - 1, 10, 1),
        "harvest_date": today,
        "wheat_variety": "", "row_width": "", "seeding_rate": "",
        "fall_n": "0", "fall_p": "0", "fall_k": "0", "fall_other": "",
        "ws_n1_date": datetime.date(today.year, 3, 1), "ws_n1_rate": "0",
        "ws_n2_date": datetime.date(today.year, 4, 1), "ws_n2_rate": "0",
        "manure_used": "No", "manure_type": "", "manure_tons": "0",
        "manure_date": today,
        "growth_reg": "", "fall_pest": "", "spring_pest": "",
        "heading_pest": "", "biologicals": "", "tillage_used": "",
        "h_length": 0.0, "h_width": 0.0,
        "moisture_readings": [], "_moisture_gen": 0,
        "test_weight": 60.0, "grain_weight": 0.0,
        "agent_notes": "",
        "h_area_ft2": 0.0, "h_acres": 0.0, "gm_avg": 0.0,
        "official_yield": 0.0, "_agree_gen": 0,
        "_gd_error": None,
    }

def _init_state():
    for k, v in _blank_defaults().items():
        if k not in st.session_state:
            st.session_state[k] = v
    if "session_entries" not in st.session_state:
        st.session_state["session_entries"] = []

def _clear_form():
    for k, v in _blank_defaults().items():
        st.session_state[k] = v
    st.rerun()

def _recompute():
    length = float(st.session_state.get("h_length", 0.0) or 0.0)
    width  = float(st.session_state.get("h_width",  0.0) or 0.0)
    ft2    = length * width
    acres  = ft2 / 43560.0
    st.session_state["h_area_ft2"] = round(ft2, 1)
    st.session_state["h_acres"]    = round(acres, 4)
    readings = [r for r in st.session_state.get("moisture_readings", []) if r > 0]
    gm_avg   = sum(readings) / len(readings) if readings else 0.0
    st.session_state["gm_avg"] = round(gm_avg, 2)
    gw  = float(st.session_state.get("grain_weight", 0.0) or 0.0)
    yld = gw * ((100 - gm_avg) / 86.5) / 60 / acres if (gw > 0 and gm_avg > 0 and acres > 0) else 0.0
    st.session_state["official_yield"] = round(yld, 2)

def _add_reading_cb():
    gen = st.session_state.get("_moisture_gen", 0)
    val = float(st.session_state.get(f"moisture_input_{gen}", 0.0) or 0.0)
    if val > 0:
        readings = st.session_state.get("moisture_readings", [])
        if len(readings) < 3:
            readings.append(round(val, 1))
            st.session_state["moisture_readings"] = readings
            st.session_state["_moisture_gen"] = gen + 1
            _recompute()
    else:
        st.session_state["_moisture_warn"] = True

def _clear_readings_cb():
    st.session_state["moisture_readings"] = []
    st.session_state["_moisture_gen"] = st.session_state.get("_moisture_gen", 0) + 1
    _recompute()

# ─────────────────────────────────────────────────────────
# EXCEL BUILDER
# ─────────────────────────────────────────────────────────

def _thin_border():
    s = Side(style="thin", color="AAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)

def _hdr(ws, row, col, text, bg="2E4057", fg="FFFFFF", bold=True, size=10):
    c = ws.cell(row=row, column=col, value=text)
    c.font      = Font(bold=bold, color=fg, name="Arial", size=size)
    c.fill      = PatternFill("solid", fgColor=bg)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border    = _thin_border()
    return c

def _dat(ws, row, col, value="", bg="FFFFFF"):
    c = ws.cell(row=row, column=col, value=value)
    c.font      = Font(name="Arial", size=10)
    c.fill      = PatternFill("solid", fgColor=bg)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    c.border    = _thin_border()
    return c

def build_excel_with_entry(data: dict, filepath: str) -> int:
    area = COUNTY_AREA.get(data.get("County",""), 4)
    data["Area"] = area
    if Path(filepath).exists():
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active
        try:
            last_id = int(ws.cell(row=ws.max_row, column=1).value or 0)
        except Exception:
            last_id = ws.max_row - 3
        entry_id = last_id + 1
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Wheat Contest Entries"
        ws.merge_cells(f"A1:{get_column_letter(len(COLUMNS))}1")
        t = ws["A1"]
        t.value     = f"Kentucky Wheat Yield Contest {CURRENT_YEAR} - Master Entry Database"
        t.font      = Font(bold=True, size=14, color="FFFFFF", name="Arial")
        t.fill      = PatternFill("solid", fgColor=HEADER_FILL)
        t.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28
        for label, c1, c2, color in SECTION_SPANS:
            ws.merge_cells(start_row=2, start_column=c1, end_row=2, end_column=c2)
            _hdr(ws, 2, c1, label, bg=color, size=9)
        ws.row_dimensions[2].height = 18
        for ci, col in enumerate(COLUMNS, 1):
            _hdr(ws, 3, ci, col.replace("_"," "), bg="4A4A4A", size=9)
        ws.row_dimensions[3].height = 42
        for i in range(1, len(COLUMNS)+1):
            ws.column_dimensions[get_column_letter(i)].width = 14
        ws.freeze_panes = "A4"
        entry_id = 1
    data["Entry_ID"]        = entry_id
    data["Submission_Date"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    row_num = ws.max_row + 1
    bg = "F7F9FC" if entry_id % 2 == 0 else "FFFFFF"
    for ci, col in enumerate(COLUMNS, 1):
        _dat(ws, row_num, ci, data.get(col,""), bg=bg)
    wb.save(filepath)
    return entry_id

# ─────────────────────────────────────────────────────────
# MAIN APP
# ─────────────────────────────────────────────────────────

def main():
    _init_state()
    st.set_page_config(page_title=f"KY Wheat Contest {CURRENT_YEAR}",
                       page_icon="🌾", layout="wide",
                       initial_sidebar_state="expanded")

    st.markdown("""
    <style>
    .main{background-color:#f0f4f8}
    .hero-banner{background:linear-gradient(135deg,#1a3a1a 0%,#2d5a27 40%,#4a7c3f 70%,#1F4E79 100%);
        border-radius:14px;padding:32px 40px 28px 40px;margin-bottom:24px;
        position:relative;overflow:hidden;box-shadow:0 6px 24px rgba(0,0,0,0.18)}
    .hero-banner::before{content:"🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾";
        position:absolute;top:8px;left:0;right:0;font-size:1.4rem;opacity:0.12;
        letter-spacing:6px;white-space:nowrap;overflow:hidden}
    .hero-banner::after{content:"🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾🌾";
        position:absolute;bottom:8px;left:0;right:0;font-size:1.4rem;opacity:0.12;
        letter-spacing:6px;white-space:nowrap;overflow:hidden}
    .hero-title{font-size:2.4rem;font-weight:800;color:#fff;
        text-shadow:0 2px 8px rgba(0,0,0,0.4);margin:0;line-height:1.15;letter-spacing:-0.5px}
    .hero-subtitle{font-size:1.05rem;color:#c8e6c9;margin-top:6px;font-weight:400;letter-spacing:0.3px}
    .hero-year{font-size:3.2rem;font-weight:900;color:rgba(255,255,255,0.18);
        position:absolute;right:40px;top:50%;transform:translateY(-50%);
        font-family:Georgia,serif;letter-spacing:-2px}
    .hero-badge{display:inline-block;background:rgba(255,255,255,0.15);
        border:1px solid rgba(255,255,255,0.3);border-radius:20px;padding:3px 14px;
        font-size:0.82rem;color:#e8f5e9;margin-top:10px;backdrop-filter:blur(4px)}
    .sec-hdr{font-size:1.05rem;font-weight:700;padding:7px 14px;
        border-radius:5px;margin:20px 0 8px 0;color:white}
    .s1{background:#1F4E79}.s1b{background:#2E5FA3}.s2{background:#375623}
    .s3{background:#7B3F00}.s4{background:#6B2737}.s5{background:#4A235A}
    .s6{background:#7E5109}.s7{background:#1A5276}.s8{background:#555555}
    /* Subsection divider inside Section 1 */
    .sub-divider{border:none;border-top:2px dashed #cbd5e0;margin:18px 0 14px 0}
    .sub-label{font-size:0.8rem;font-weight:700;text-transform:uppercase;
        letter-spacing:0.8px;color:#718096;margin-bottom:8px;margin-top:4px}
    .metric-box{background:#f8f9fa;border:1px solid #dee2e6;
        border-radius:6px;padding:10px 14px;margin-top:4px}
    .metric-label{font-size:0.78rem;color:#6c757d;margin-bottom:2px}
    .metric-value{font-size:1.5rem;font-weight:700;color:#212529}
    .metric-ok{color:#198754}.metric-warn{color:#dc3545}
    .moisture-badge{background:#e9ecef;border-radius:4px;padding:3px 9px;
        font-size:0.88rem;display:inline-block;margin:2px 3px}
    .agreement-box{background:#fff3cd;border:2px solid #ffc107;
        border-radius:8px;padding:16px 20px;margin:24px 0 8px 0}
    .contact-footer{background:#f8f9fa;border:1px solid #dee2e6;border-radius:10px;
        padding:16px 20px;text-align:center;margin-top:30px;font-size:0.9rem;color:#555}
    .contact-footer a{color:#1F4E79;font-weight:600;text-decoration:none}
    </style>
    """, unsafe_allow_html=True)

    # ── Hero header ──────────────────────────────────────
    st.markdown(f"""
    <div class="hero-banner">
      <div class="hero-year">{CURRENT_YEAR}</div>
      <div class="hero-title">🌾 Kentucky Wheat Yield Contest</div>
      <div class="hero-subtitle">
        University of Kentucky Cooperative Extension — Digital Entry Form
      </div>
      <div class="hero-badge">📅 Contest Year {CURRENT_YEAR} &nbsp;·&nbsp; Official Submission Portal</div>
    </div>
    """, unsafe_allow_html=True)

    col_clr, _ = st.columns([1, 5])
    with col_clr:
        if st.button("🔄 Clear Form", use_container_width=True):
            _clear_form()
    st.divider()

    # ── Sidebar ──────────────────────────────────────────
    with st.sidebar:
        st.markdown(f"### 🌾 KY Wheat Contest {CURRENT_YEAR}")
        st.divider()
        st.markdown("**📄 Contest Rules**")
        rules_path = Path("2025WheatYieldContestRules.pdf")
        if rules_path.exists():
            with open(rules_path, "rb") as f:
                st.download_button("⬇️ Download Contest Rules (PDF)",
                    data=f, file_name="KY_Wheat_Contest_Rules.pdf",
                    mime="application/pdf", use_container_width=True)
        else:
            st.info(
                "- Min. **1.5 acres** harvested\n"
                "- Deadline: **July 31**\n"
                "- Supervisor must witness harvest\n"
                "- Submit grain sample to **Colette Laurent**, Princeton KY\n"
                "- Official yield at **13.5% moisture**\n\n"
                "_Place `2025WheatYieldContestRules.pdf` in app folder to enable PDF download._"
            )
        st.divider()
        cfg = _gdrive_secrets()
        if all([cfg.get("client_email"), cfg.get("private_key"), cfg.get("sheet_id")]):
            st.success("📊 Google Sheets: connected")
        else:
            st.warning("📊 Google Sheets: not configured")

    # ════════════════════════════════════════════════════
    # SECTION 1 — PRODUCER INFO  (grouped clearly)
    # ════════════════════════════════════════════════════
    st.markdown('<div class="sec-hdr s1">👤 Section 1 — Producer & Supervisor Information</div>',
                unsafe_allow_html=True)

    # ── Producer block ───────────────────────────────────
    st.markdown('<div class="sub-label">🌿 Producer / Grower</div>', unsafe_allow_html=True)
    p1, p2, p3 = st.columns(3)
    with p1:
        st.selectbox("County *", ["— Select —"] + KY_COUNTIES, key="county")
        st.text_input("Producer Full Name *", key="producer_name")
        st.text_input("Profession / Operation Type *", key="profession")
    with p2:
        st.text_input("Producer Email *", key="producer_email",
                      placeholder="grower@email.com")
        st.text_input("Phone *", key="producer_phone",
                      placeholder="270-555-1234")
        st.text_input("Mobile", key="producer_mobile")
    with p3:
        st.text_input("Street Address", key="producer_address")
        st.text_input("Town / City", key="producer_town")
        st.text_input("Zip Code", key="producer_zip")

    # ── Supervisor block ─────────────────────────────────
    st.markdown('<hr class="sub-divider"><div class="sub-label">🏛️ County Agent / Supervisor</div>',
                unsafe_allow_html=True)
    s1, s2, s3 = st.columns(3)
    with s1:
        st.text_input("Supervisor Full Name *", key="supervisor_name")
    with s2:
        st.text_input("Supervisor Phone *", key="supervisor_phone",
                      placeholder="270-555-5678")
    with s3:
        st.date_input("Supervisor Signature Date", key="supervisor_sig_date")

    # ════════════════════════════════════════════════════
    # SECTION 2 — AGRONOMIC  (planting + harvest dates here)
    # ════════════════════════════════════════════════════
    st.markdown('<div class="sec-hdr s2">🌱 Section 2 — Agronomic Data</div>',
                unsafe_allow_html=True)
    a1, a2, a3 = st.columns(3)
    with a1:
        st.radio("Contest Division *",
                 ["Division I - Tillage (conv./min.)", "Division II - No-Tillage"],
                 horizontal=True, key="division")
        st.selectbox("Previous Crop", ["Corn","Soybeans","Other"], key="previous_crop")
        st.text_input("Wheat Variety *", key="wheat_variety")
    with a2:
        st.date_input("Planting Date *", key="planting_date")
        st.date_input("Harvest Date *", key="harvest_date")
    with a3:
        st.text_input("Row Width (inches)", key="row_width")
        st.text_input("Seeding Rate (seeds/A or lb/A)", key="seeding_rate")

    # ════════════════════════════════════════════════════
    # SECTION 3 — FERTILIZER
    # ════════════════════════════════════════════════════
    st.markdown('<div class="sec-hdr s3">🧪 Section 3 — Fertilizer</div>',
                unsafe_allow_html=True)
    st.caption("Fall Fertilizer (lb/acre)")
    c1,c2,c3,c4 = st.columns(4)
    with c1: st.text_input("N (lb/A)", key="fall_n")
    with c2: st.text_input("P2O5 (lb/A)", key="fall_p")
    with c3: st.text_input("K2O (lb/A)", key="fall_k")
    with c4: st.text_input("Other fertilizers", key="fall_other")
    st.caption("Winter/Spring Nitrogen")
    c1,c2,c3,c4 = st.columns(4)
    with c1: st.date_input("Application 1 Date", key="ws_n1_date")
    with c2: st.text_input("N lb/A (App 1)", key="ws_n1_rate")
    with c3: st.date_input("Application 2 Date", key="ws_n2_date")
    with c4: st.text_input("N lb/A (App 2)", key="ws_n2_rate")
    st.caption("Manure")
    c1,c2,c3,c4 = st.columns(4)
    with c1: st.selectbox("Manure (last 18 months)?", ["No","Yes"], key="manure_used")
    manure_on = st.session_state.get("manure_used","No") == "Yes"
    with c2: st.text_input("Type", key="manure_type", disabled=not manure_on)
    with c3: st.text_input("Tons/A", key="manure_tons", disabled=not manure_on)
    with c4: st.date_input("Manure Date", key="manure_date", disabled=not manure_on)

    # ════════════════════════════════════════════════════
    # SECTION 4 — PEST MANAGEMENT
    # ════════════════════════════════════════════════════
    st.markdown('<div class="sec-hdr s4">🛡️ Section 4 — Pest Management & Other Inputs</div>',
                unsafe_allow_html=True)
    c1, c2 = st.columns(2)
    with c1:
        st.text_area("Growth Regulator(s)", height=68, key="growth_reg")
        st.text_area("Fall Pest Products", height=68, key="fall_pest")
        st.text_area("Biologicals / Other", height=68, key="biologicals")
    with c2:
        st.text_area("Spring Pest Products", height=68, key="spring_pest")
        st.text_area("Heading/Flowering Pest Products", height=68, key="heading_pest")
        st.text_input("Tillage Equipment / Method Used", key="tillage_used")

    # ════════════════════════════════════════════════════
    # SECTION 5 — HARVEST AREA
    # ════════════════════════════════════════════════════
    st.markdown('<div class="sec-hdr s5">📐 Section 5 — Harvest Area Measurement</div>',
                unsafe_allow_html=True)
    st.info("Minimum harvest area: **1.50 acres** (65,340 sq ft).")
    c1,c2,c3,c4 = st.columns(4)
    with c1:
        st.number_input("Length (ft) *", min_value=0.0, step=1.0, format="%.1f",
                        key="h_length", on_change=_recompute)
    with c2:
        st.number_input("Width (ft) *", min_value=0.0, step=1.0, format="%.1f",
                        key="h_width", on_change=_recompute)
    h_area_ft2 = st.session_state["h_area_ft2"]
    h_acres    = st.session_state["h_acres"]
    area_ok    = h_acres >= 1.50
    with c3:
        st.markdown(f'<div class="metric-box"><div class="metric-label">Area (ft²)</div>'
                    f'<div class="metric-value">{h_area_ft2:,.1f}</div></div>',
                    unsafe_allow_html=True)
    with c4:
        cls = "metric-ok" if area_ok else "metric-warn"
        st.markdown(f'<div class="metric-box"><div class="metric-label">Acres — '
                    f'{"OK ✅" if area_ok else "Below min ❌"}</div>'
                    f'<div class="metric-value {cls}">{h_acres:.2f}</div></div>',
                    unsafe_allow_html=True)

    # ════════════════════════════════════════════════════
    # SECTION 6 — GRAIN CHARACTERISTICS
    # ════════════════════════════════════════════════════
    st.markdown('<div class="sec-hdr s6">🌡️ Section 6 — Grain Characteristics</div>',
                unsafe_allow_html=True)
    readings_now = st.session_state.get("moisture_readings", [])
    n_done       = len(readings_now)
    slots_left   = 3 - n_done
    warn_empty   = st.session_state.pop("_moisture_warn", False)
    st.caption("**Grain Moisture:** Certified tester → 1 reading. Handheld → 3 readings.")
    mc1,mc2,mc3 = st.columns([2,1,3])
    with mc1:
        gen   = st.session_state.get("_moisture_gen", 0)
        label = (f"Moisture Reading {n_done+1} of 3 (%)"
                 if slots_left > 0 else "All 3 readings recorded")
        st.number_input(label, min_value=0.0, max_value=40.0, step=0.1,
                        format="%.1f", value=0.0, key=f"moisture_input_{gen}",
                        disabled=(slots_left==0))
        st.session_state["moisture_input"] = st.session_state.get(f"moisture_input_{gen}", 0.0)
        if warn_empty:
            st.warning("Enter a value > 0 before adding.")
    with mc2:
        st.markdown("<div style='margin-top:28px'></div>", unsafe_allow_html=True)
        st.button("➕ Add Reading", disabled=(slots_left==0),
                  use_container_width=True, on_click=_add_reading_cb)
        if readings_now:
            st.button("🗑️ Clear", use_container_width=True, on_click=_clear_readings_cb)
    with mc3:
        gm_avg = st.session_state["gm_avg"]
        if readings_now:
            badges = " ".join(
                f'<span class="moisture-badge">#{i+1}: <b>{r}%</b></span>'
                for i, r in enumerate(readings_now))
            st.markdown(f'<div class="metric-box"><div class="metric-label">'
                        f'Recorded readings &nbsp; {badges}</div>'
                        f'<div class="metric-value metric-ok">Avg: {gm_avg:.1f}%</div></div>',
                        unsafe_allow_html=True)
        else:
            st.markdown('<div class="metric-box"><div class="metric-label">Recorded readings</div>'
                        '<div class="metric-value" style="color:#adb5bd;font-size:1rem">'
                        'None yet &mdash; add up to 3</div></div>', unsafe_allow_html=True)
    st.markdown("")
    st.number_input("Test Weight (lb/bu)", min_value=0.0, max_value=70.0,
                    step=0.1, format="%.1f", key="test_weight")

    # ════════════════════════════════════════════════════
    # SECTION 7 — YIELD
    # ════════════════════════════════════════════════════
    st.markdown('<div class="sec-hdr s7">📊 Section 7 — Yield Calculation</div>',
                unsafe_allow_html=True)
    st.caption("Formula: lbs x [(100 - %moisture) / 86.5] / 60 lb/bu / acres")
    c1,c2,c3 = st.columns(3)
    with c1:
        st.number_input("Grain Weight from Scale (lbs) *", min_value=0.0,
                        step=10.0, format="%.1f", key="grain_weight", on_change=_recompute)
    with c2:
        official_yield = st.session_state["official_yield"]
        yld_cls = "metric-ok" if official_yield > 0 else ""
        st.markdown(f'<div class="metric-box"><div class="metric-label">Official Yield (Bu/Acre)</div>'
                    f'<div class="metric-value {yld_cls}">{official_yield:.2f}</div></div>',
                    unsafe_allow_html=True)
    with c3:
        gw = float(st.session_state.get("grain_weight", 0.0) or 0.0)
        st.markdown(f'<div class="metric-box" style="font-size:0.82rem;color:#495057;line-height:1.7">'
                    f'<b>Step-by-step:</b><br>{gw:.0f} x [(100-{gm_avg:.1f})/86.5]'
                    f'<br>&divide;60 &divide;{h_acres:.2f}ac'
                    f'<br>= <b>{official_yield:.2f} bu/ac</b></div>', unsafe_allow_html=True)

    # ════════════════════════════════════════════════════
    # SECTION 8 — NOTES + SCALE TICKET PHOTO
    # ════════════════════════════════════════════════════
    st.markdown('<div class="sec-hdr s8">📝 Section 8 — Notes & Scale Ticket Photo</div>',
                unsafe_allow_html=True)
    col_notes, col_photo = st.columns([1.1, 0.9], gap="large")
    with col_notes:
        st.markdown("""
        <div style="background:#f8f9fa;border:1px solid #dee2e6;border-radius:10px;
                    padding:18px 20px 6px 20px;margin-bottom:8px">
          <div style="font-size:0.95rem;font-weight:700;color:#343a40;margin-bottom:6px">
            📋 Agent Notes <span style="font-weight:400;color:#6c757d;font-size:0.85rem">(optional)</span>
          </div>
          <div style="font-size:0.82rem;color:#6c757d;margin-bottom:10px">
            Record any unusual field conditions, equipment issues, weather events,
            or other observations relevant to this entry.
          </div>
        </div>
        """, unsafe_allow_html=True)
        st.text_area("Agent notes", height=130, key="agent_notes",
                     placeholder="e.g. Field had minor flooding in NE corner. "
                                 "Scale certified Oct 2025...",
                     label_visibility="collapsed")
    with col_photo:
        st.markdown("""
        <div style="background:linear-gradient(135deg,#fff8e1,#fff3cd);
                    border:2px dashed #ffc107;border-radius:10px;
                    padding:18px 20px 10px 20px;margin-bottom:8px">
          <div style="font-size:0.95rem;font-weight:700;color:#856404;margin-bottom:4px">
            📷 Scale Ticket Photo
            <span style="font-weight:400;font-size:0.85rem">(optional but recommended)</span>
          </div>
          <div style="font-size:0.82rem;color:#6c757d;margin-bottom:10px;line-height:1.5">
            Upload a photo of the <b>certified scale ticket</b> showing grain weight.
            Accepted: JPG, PNG, HEIC &nbsp;·&nbsp; Saved with your entry automatically.
          </div>
        </div>
        """, unsafe_allow_html=True)
        scale_photo = st.file_uploader("Upload scale ticket photo",
                                       type=["jpg","jpeg","png","heic","webp"],
                                       label_visibility="collapsed")
        if scale_photo:
            st.image(scale_photo,
                     caption=f"✅ {scale_photo.name}  ({scale_photo.size/1024:.0f} KB)",
                     use_container_width=True)
            st.success("Photo will be saved with this entry.")
        else:
            st.markdown("""
            <div style="text-align:center;padding:20px 0;color:#adb5bd;font-size:2rem">
              📄<div style="font-size:0.82rem;margin-top:4px;color:#ced4da">
              Drag & drop or click above to upload</div>
            </div>""", unsafe_allow_html=True)

    # ════════════════════════════════════════════════════
    # CERTIFICATION + SUBMIT
    # ════════════════════════════════════════════════════
    st.divider()
    st.markdown("""
    <div class="agreement-box">
    <b>📋 Agent Certification — Required Before Submission</b><br>
    By checking the box below, you certify that:
    <ul style="margin:8px 0 4px 20px;">
      <li>All information entered is <b>accurate and complete</b>.</li>
      <li>The harvest area was <b>physically measured</b> and meets the 1.50-acre minimum.</li>
      <li>Grain moisture and weight were recorded using <b>certified or approved equipment</b>.</li>
      <li>You witnessed or verified the harvest as the <b>supervising county agent</b>.</li>
      <li>You understand this entry will be <b>submitted officially</b> to the UK Extension state office.</li>
    </ul>
    </div>
    """, unsafe_allow_html=True)

    agree_gen = st.session_state.get("_agree_gen", 0)
    st.checkbox("I certify that all information above is accurate and complete.",
                key=f"agreement_checked_{agree_gen}", value=False)
    agreement = st.session_state.get(f"agreement_checked_{agree_gen}", False)

    errors = []
    if st.session_state.get("county","— Select —") == "— Select —":
        errors.append("County not selected")
    if not st.session_state.get("producer_name","").strip():
        errors.append("Producer name missing")
    if not st.session_state.get("producer_email","").strip():
        errors.append("Producer email missing")
    if not st.session_state.get("producer_phone","").strip():
        errors.append("Producer phone missing (required)")
    if not st.session_state.get("profession","").strip():
        errors.append("Profession / operation type missing")
    if not st.session_state.get("supervisor_name","").strip():
        errors.append("Supervisor name missing")
    if not st.session_state.get("supervisor_phone","").strip():
        errors.append("Supervisor phone missing")
    if not st.session_state.get("wheat_variety","").strip():
        errors.append("Wheat variety missing")
    if st.session_state["h_acres"] < 1.50:
        errors.append(f"Harvest area {st.session_state['h_acres']:.2f} ac < 1.50 ac minimum")
    if float(st.session_state.get("grain_weight",0.0) or 0.0) <= 0:
        errors.append("Grain weight not entered")
    if not st.session_state.get("moisture_readings"):
        errors.append("No moisture reading — press Add Reading at least once")

    if errors:
        with st.expander("⚠️ Required fields not complete", expanded=True):
            for e in errors:
                st.markdown(f"- 🔴 {e}")

    form_ready = agreement and len(errors) == 0
    if not form_ready:
        st.caption("🔒 Check the certification box above to unlock." if not agreement
                   else "🔒 Fix the highlighted fields above to unlock.")

    submit_clicked = st.button("💾  Save Entry & Send to State Office",
                               disabled=not form_ready, type="primary",
                               use_container_width=True)

    # ════════════════════════════════════════════════════
    # ON SUBMIT
    # ════════════════════════════════════════════════════
    if submit_clicked and form_ready:
        r   = st.session_state.get("moisture_readings", [])
        gm1 = r[0] if len(r) > 0 else ""
        gm2 = r[1] if len(r) > 1 else ""
        gm3 = r[2] if len(r) > 2 else ""
        photo_b64 = ""
        if scale_photo:
            photo_link = upload_photo_to_gdrive(
                scale_photo.getvalue(), scale_photo.name,
                entry_id, st.session_state["county"]
            )
        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
        data = {
            "County":                    st.session_state["county"],
            "Producer_Name":             st.session_state["producer_name"],
            "Producer_Email":            st.session_state.get("producer_email",""),
            "Producer_Phone":            st.session_state.get("producer_phone",""),
            "Producer_Mobile":           st.session_state.get("producer_mobile",""),
            "Producer_Address":          st.session_state.get("producer_address",""),
            "Producer_Town":             st.session_state.get("producer_town",""),
            "Producer_Zip":              st.session_state.get("producer_zip",""),
            "Profession":                st.session_state.get("profession",""),
            "Supervisor_Name":           st.session_state["supervisor_name"],
            "Supervisor_Phone":          st.session_state.get("supervisor_phone",""),
            "Supervisor_Signature_Date": str(st.session_state.get("supervisor_sig_date","")),
            "Division":                  st.session_state.get("division",""),
            "Previous_Crop":             st.session_state.get("previous_crop",""),
            "Planting_Date":             str(st.session_state.get("planting_date","")),
            "Harvest_Date":              str(st.session_state.get("harvest_date","")),
            "Wheat_Variety":             st.session_state["wheat_variety"],
            "Row_Width_inches":          st.session_state.get("row_width",""),
            "Seeding_Rate":              st.session_state.get("seeding_rate",""),
            "Fall_N_lbA":                st.session_state.get("fall_n","0"),
            "Fall_P2O5_lbA":             st.session_state.get("fall_p","0"),
            "Fall_K2O_lbA":              st.session_state.get("fall_k","0"),
            "Fall_Other_Fertilizer":     st.session_state.get("fall_other",""),
            "Winter_Spring_N1_Date":     str(st.session_state.get("ws_n1_date","")),
            "Winter_Spring_N1_lbA":      st.session_state.get("ws_n1_rate","0"),
            "Winter_Spring_N2_Date":     str(st.session_state.get("ws_n2_date","")),
            "Winter_Spring_N2_lbA":      st.session_state.get("ws_n2_rate","0"),
            "Manure_Used":               st.session_state.get("manure_used","No"),
            "Manure_Type":               st.session_state.get("manure_type","") if manure_on else "",
            "Manure_TonsA":              st.session_state.get("manure_tons","") if manure_on else "",
            "Manure_Date":               str(st.session_state.get("manure_date","")) if manure_on else "",
            "Growth_Regulator":          st.session_state.get("growth_reg",""),
            "Fall_Pest_Products":        st.session_state.get("fall_pest",""),
            "Spring_Pest_Products":      st.session_state.get("spring_pest",""),
            "Heading_Flowering_Pest":    st.session_state.get("heading_pest",""),
            "Biologicals_Other":         st.session_state.get("biologicals",""),
            "Tillage_Used":              st.session_state.get("tillage_used",""),
            "Harvest_Length_ft":         st.session_state["h_length"],
            "Harvest_Width_ft":          st.session_state["h_width"],
            "Harvest_Area_ft2":          st.session_state["h_area_ft2"],
            "Harvest_Acres":             st.session_state["h_acres"],
            "Grain_Moisture_1":          gm1, "Grain_Moisture_2": gm2, "Grain_Moisture_3": gm3,
            "Grain_Moisture_Avg":        st.session_state["gm_avg"],
            "Test_Weight_lbbu":          st.session_state.get("test_weight",60.0),
            "Grain_Weight_lbs":          st.session_state["grain_weight"],
            "Official_Yield_BuAcre":     st.session_state["official_yield"],
            "Agent_Notes":               st.session_state.get("agent_notes",""),
            "Scale_Ticket_Photo":        photo_link if photo_link else "[no photo]",
            "_moisture_list":            r,
            "Submission_Date":           now_str,
        }
        try:
            entry_id = build_excel_with_entry(data, EXCEL_FILE)
            area     = COUNTY_AREA.get(data["County"], 4)
            st.session_state["_gd_error"] = None
            sheet_ok = append_entry_to_sheet(data, entry_id)
            subject  = (f"KY Wheat Contest Entry #{entry_id} — "
                        f"{data['County']} County — {data['Producer_Name']}")
            email_ok = send_formsubmit_email(data, entry_id, subject)

            st.success(f"✅ Entry #{entry_id} submitted successfully!")
            st.markdown(
                f"**Producer:** {data['Producer_Name']} &nbsp;|&nbsp; "
                f"**County:** {data['County']} (Area {area}) &nbsp;|&nbsp; "
                f"**Division:** {data['Division'].split('-')[0].strip()}  \n"
                f"**Official Yield:** {data['Official_Yield_BuAcre']:.2f} bu/acre &nbsp;|&nbsp; "
                f"**Harvest Area:** {data['Harvest_Acres']:.2f} acres &nbsp;|&nbsp; "
                f"**Grain Moisture:** {data['Grain_Moisture_Avg']:.1f}%"
            )
            col1, col2, col3 = st.columns(3)
            with col1: st.success("📊 Saved to Excel")
            with col2:
                if sheet_ok:
                    st.success("☁️ Synced to Google Sheets")
                else:
                    st.warning("☁️ Sheets sync failed")
                    if st.session_state.get("_gd_error"):
                        with st.expander("Error detail"):
                            st.code(st.session_state["_gd_error"])
            with col3:
                if email_ok:
                    st.success("📧 Email notification sent")
                else:
                    st.warning(f"📧 Email failed")

            with open(EXCEL_FILE, "rb") as f:
                st.download_button("⬇️ Download My Entry (Excel backup)", data=f,
                    file_name=f"entry_{entry_id}_{data['County']}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    help="Download a local backup. This does NOT send another email.")

            st.session_state["session_entries"].append({
                "Entry #": entry_id, "Producer": data["Producer_Name"],
                "County": data["County"],
                "Yield Bu/A": f"{data['Official_Yield_BuAcre']:.2f}", "Time": now_str,
            })
            st.session_state["_agree_gen"] = st.session_state.get("_agree_gen", 0) + 1

        except Exception as ex:
            st.error(f"Error saving entry: {ex}")

    # ── Session entries ──────────────────────────────────
    session_entries = st.session_state.get("session_entries", [])
    if session_entries:
        st.divider()
        st.subheader(f"📋 Your Submissions This Session ({len(session_entries)})")
        st.caption("Entries submitted during this browser session — securely saved to the state office.")
        st.dataframe(pd.DataFrame(session_entries), use_container_width=True, hide_index=True)

    # ── Contact footer ───────────────────────────────────
    st.divider()
    st.markdown(f"""
    <div class="contact-footer">
        <b>Questions or issues with this form?</b><br>
        Contact Dr. Chad Lee, Extension Professor, UK:
        <a href="mailto:{CONTACT_EMAIL}">✉️ {CONTACT_EMAIL}</a>
        &nbsp;&nbsp;|&nbsp;&nbsp;
        <span style="color:#888;font-size:0.82rem">
        Kentucky Wheat Yield Contest {CURRENT_YEAR} &mdash;
        University of Kentucky Cooperative Extension
        </span>
    </div>
    """, unsafe_allow_html=True)


if __name__ == "__main__":
    main()
